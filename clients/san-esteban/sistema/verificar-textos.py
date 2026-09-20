#!/usr/bin/env python3
"""Compara los textos de las piezas contra el brief, celda por celda.

    python3 clients/san-esteban/sistema/verificar-textos.py <brief.xlsx>

Existe porque «texto verbatim del brief» es la regla que más caro sale romper y a
ojo no se verifica: hay que leer la celda «TEXTO SOBRE LA IMAGEN» y compararla con
lo que quedó en la pieza. Reporta tres estados:

  ✓ literal                     coincide carácter por carácter
  ⚠ excepción acordada          la única desviación autorizada (100 → 110 años,
                                decidida el 07-09-2026)
  ✖ DIFIERE                     cualquier otra cosa: hay que arreglarla

Los reels se verifican contra src/compositions/sanEstebanReelesOctubre.ts.
"""
import re
import sys
import pathlib

import openpyxl

AQUI = pathlib.Path(__file__).parent
sys.path.insert(0, str(AQUI))
import build  # noqa: E402


def norm(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def main(ruta_brief):
    ws = openpyxl.load_workbook(ruta_brief, data_only=True)["Brief"]
    piezas = {}
    for r in range(10, 30):
        n = ws.cell(r, 1).value
        if not n:
            continue
        d = {}
        for linea in (ws.cell(r, 7).value or "").split("\n"):
            for k in ("TÍTULO", "BAJADA", "APOYO"):
                if linea.startswith(k + ":"):
                    d[k] = linea.split(":", 1)[1].strip()
        piezas[str(n).zfill(2)] = d

    usados = {p["id"][1:]: p for p in build.PIEZAS}
    ok = exc = dif = 0
    for num, d in piezas.items():
        if num not in usados:
            print(f"{num}  — reel, se verifica en sanEstebanReelesOctubre.ts")
            continue
        p = usados[num]
        for campo, valor in [("TÍTULO", p["titular"]),
                             ("BAJADA", p["bajada"]),
                             ("APOYO", p["cta"])]:
            esperado, real = norm(d.get(campo)), norm(valor)
            if esperado == real:
                print(f"{num} {campo:7} ✓ literal")
                ok += 1
                continue
            e2 = norm(re.sub(r"[Mm]ás de 100 años", "110 años", esperado))
            if e2 == real:
                print(f"{num} {campo:7} ⚠ excepción acordada «100 → 110 años»")
                exc += 1
            else:
                print(f"{num} {campo:7} ✖ DIFIERE\n     brief: {esperado}\n     pieza: {real}")
                dif += 1
    print(f"\n{ok} literales · {exc} con la excepción acordada · {dif} sin justificar")
    return 1 if dif else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("uso: verificar-textos.py <brief.xlsx>")
    sys.exit(main(sys.argv[1]))
