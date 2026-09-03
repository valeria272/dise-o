# -*- coding: utf-8 -*-
"""Instantánea de texto de una grilla de Hilton: es la BASE para el diff de la
ronda siguiente. El .xlsx pesa decenas de MB y no viaja en git; esto sí.

Uso:
    PYTHONIOENCODING=utf-8 PYTHONUTF8=1 python scripts/grilla-instantanea.py \
        raw/hilton/dt/DT-grilla-septiembre-2026.xlsx \
        clients/hilton/grillas/dt-septiembre-2026.md DOUBLETREE

Lee los tachados (`font.strike`) y los hipervínculos, que es justo lo que se
pierde al copiar y pegar la celda a mano."""
import sys, openpyxl, datetime
from openpyxl.cell.rich_text import CellRichText

SRC   = sys.argv[1]
OUT   = sys.argv[2]
MARCA = sys.argv[3] if len(sys.argv) > 3 else "?"

wb   = openpyxl.load_workbook(SRC, data_only=True)
wbr  = openpyxl.load_workbook(SRC, rich_text=True)   # para leer los tachados
wbl  = openpyxl.load_workbook(SRC)                   # para leer los hipervínculos

def tachado(hoja, coord):
    v = wbr[hoja][coord].value
    if not isinstance(v, CellRichText):
        return []
    return [getattr(r, "text", str(r)) for r in v
            if getattr(r, "font", None) and r.font.strike]

def enlace(hoja, coord):
    h = wbl[hoja][coord].hyperlink
    return h.target if h else None

L = []
L.append(f"# Grilla {MARCA} — instantánea")
L.append("")
L.append(f"> Instantánea del {datetime.date.today():%d-%m-%Y} tomada con `/al-dia`.")
L.append(f"> Origen: `{SRC}`.")
L.append("> **Sirve de base para el DIFF de la próxima ronda**: los comentarios se")
L.append("> prependen sobre los viejos en la misma celda, así que sin esta copia no se")
L.append("> distingue lo nuevo de lo ya resuelto.")
L.append("")

for hoja, fila_estado in [("FEED", 15), ("STORIES", 16), ("REELSORGÁNICOS", 15)]:
    ws = wb[hoja]
    L.append(f"## {hoja}")
    L.append("")
    campos = {c[0].row: str(c[0].value).strip() for c in ws.iter_rows(min_col=1, max_col=1)
              if c[0].value}
    for col in range(2, ws.max_column + 1):
        letra = openpyxl.utils.get_column_letter(col)
        estado = ws.cell(row=fila_estado, column=col).value
        if not estado:
            continue
        L.append(f"### Columna {letra} — **{estado}**")
        L.append("")
        for fila, campo in campos.items():
            if fila == fila_estado:
                continue
            v = ws.cell(row=fila, column=col).value
            if v is None or not str(v).strip():
                continue
            v = str(v).strip()
            if isinstance(ws.cell(row=fila, column=col).value, datetime.datetime):
                v = f"{ws.cell(row=fila, column=col).value:%d-%m-%Y}"
            url = enlace(hoja, f"{letra}{fila}")
            L.append(f"**{campo}:**" + (f" ([enlace]({url}))" if url else ""))
            L.append("")
            L.append("```")
            L.append(v)
            L.append("```")
            tach = tachado(hoja, f"{letra}{fila}")
            if tach:
                L.append("")
                L.append("*Tachado en la grilla (ya resuelto, NO volver a aplicar):*")
                for t in tach:
                    L.append(f"> ~~{t.strip()}~~")
            L.append("")
        L.append("")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(L) + "\n")
print(f"escrito {OUT}")
