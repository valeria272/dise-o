# -*- coding: utf-8 -*-
"""Presentación comercial Más Center — Proyectos Operativos.

Rehace la PPT interna (56 láminas 4:3 en Calibri) con el sistema del brochure de
Algarrobal. Sale un .pptx editable que se sube a Google Slides y de ahí se exporta
el PDF, así los tres formatos vienen del mismo origen y no se despegan.
"""
import json, sys, urllib.parse, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
import mascenter_sistema as S

RAIZ = pathlib.Path(__file__).resolve().parent.parent
BASE = RAIZ / "raw/mascenter-presentacion"
A, AL, M = S.ANCHO, S.ALTO, S.MARGEN


def maps(c):
    q = urllib.parse.quote(f'Más Center {c["nombre"]}, {c["direccion"]}, {c["comuna"]}, Chile')
    return f"https://www.google.com/maps/search/?api=1&query={q}"


def lamina(prs, fondo=S.BLANCO):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    S.rect(s, 0, 0, A, AL, fondo)
    return s


def foto_cubriendo(slide, ruta, x, y, ancho, alto):
    """Inserta la imagen recortada al marco (cover), sin deformarla."""
    im = Image.open(ruta); pr = im.width/im.height; mr = ancho/alto
    if pr > mr:                      # sobra ancho -> recorto a los lados
        w = int(im.height*mr); im = im.crop(((im.width-w)//2, 0, (im.width-w)//2+w, im.height))
    else:                            # sobra alto -> recorto arriba y abajo
        h = int(im.width/mr); im = im.crop((0, (im.height-h)//2, im.width, (im.height-h)//2+h))
    tmp = BASE/"_tmp"; tmp.mkdir(exist_ok=True)
    d = tmp/(pathlib.Path(ruta).stem + "_cover.png"); im.save(d)
    return slide.shapes.add_picture(str(d), x, y, ancho, alto)


def foto_entera(slide, ruta, x, y, ancho, alto):
    """Mete la imagen completa dentro del marco, centrada, sin recortar."""
    im = Image.open(ruta); pr = im.width/im.height
    w, h = (ancho, int(ancho/pr)) if ancho/pr <= alto else (int(alto*pr), alto)
    return slide.shapes.add_picture(str(ruta), int(x+(ancho-w)/2), int(y+(alto-h)/2), w, h)


# ── láminas ──────────────────────────────────────────────────────────────────
def portada(prs, foto):
    s = lamina(prs)
    foto_cubriendo(s, foto, 0, 0, A, AL)
    S.rect(s, 0, 0, A, AL, S.GRAFITO, transparencia=0.28)
    S.campo(s, A, AL, 0.46, S.ROJO)
    s.shapes.add_picture(str(BASE/"assets/logo-mascenter-blanco.png"), M, int(AL*0.13), Emu(2050000))
    S.texto(s, M, int(AL*0.44), int(A*0.55), Emu(400000), "PRESENTACIÓN COMERCIAL",
            tam=13, peso="SemiBold", color=S.BLANCO, espaciado=3.2)
    S.texto(s, M, int(AL*0.52), int(A*0.58), Emu(1900000), "Proyectos\nen operación",
            tam=54, peso="Bold", color=S.BLANCO, interlineado=0.98)
    S.texto(s, M, int(AL*0.855), int(A*0.6), Emu(400000),
            "25 strip centers · 8 regiones de Chile", tam=15, peso="Light", color=S.BLANCO)
    s.shapes.add_picture(str(BASE/"assets/logo-grupoifb-blanco.png"),
                         int(A-M-Emu(950000)), int(AL-M-Emu(600000)), Emu(950000))
    return s


def resumen(prs, foto):
    s = lamina(prs)
    foto_cubriendo(s, foto, int(A*0.45), 0, int(A*0.55), AL)
    S.campo(s, A, AL, 0.47, S.ROJO)
    S.texto(s, M, int(AL*0.16), int(A*0.40), Emu(400000), "QUIÉNES SOMOS",
            tam=13, peso="SemiBold", color=S.BLANCO, espaciado=3.2)
    S.texto(s, M, int(AL*0.24), int(A*0.40), Emu(1200000),
            "Desarrollamos y operamos\nstrip centers",
            tam=31, peso="Bold", color=S.BLANCO, interlineado=1.06)
    S.texto(s, M, int(AL*0.50), int(A*0.395), Emu(2200000),
            "Más Center es la unidad de centros comerciales de Grupo IFB. "
            "Operamos una red de 25 proyectos en 8 regiones del país, anclados por "
            "supermercados y farmacias con contratos de largo plazo.\n\n"
            "Cada centro se diseña para el comercio de proximidad: alto flujo, "
            "acceso directo y un mix comercial que sostiene la rentabilidad del local.",
            tam=13.5, peso="Light", color=S.BLANCO, interlineado=1.5)
    return s


def indice(prs, secciones):
    s = lamina(prs, S.PAPEL)
    S.campo(s, A, AL, 0.012, S.ROJO, punta=0.022)   # filo rojo discreto
    S.texto(s, M, int(AL*0.16), int(A*0.5), Emu(400000), "CONTENIDO",
            tam=13, peso="SemiBold", color=S.ROJO, espaciado=3.2)
    y = int(AL*0.28)
    for i, (t, d) in enumerate(secciones, 1):
        S.texto(s, M, y, Emu(700000), Emu(700000), str(i), tam=34, peso="Bold", color=S.ROJO)
        S.texto(s, int(M+Emu(750000)), y+Emu(40000), int(A*0.62), Emu(400000), t,
                tam=20, peso="SemiBold", color=S.GRAFITO)
        S.texto(s, int(M+Emu(750000)), y+Emu(400000), int(A*0.62), Emu(400000), d,
                tam=12, peso="Light", color=S.GRAFITO)
        y += Emu(1000000)
    S.marca_pie(s)
    return s


def portadilla(prs, n, titulo, foto):
    s = lamina(prs)
    foto_cubriendo(s, foto, int(A*0.40), 0, int(A*0.60), AL)
    S.campo(s, A, AL, 0.44, S.ROJO)
    S.texto(s, M, int(AL*0.33), Emu(1200000), Emu(1400000), str(n),
            tam=96, peso="Bold", color=S.BLANCO)
    S.texto(s, M, int(AL*0.60), int(A*0.40), Emu(1200000), titulo,
            tam=30, peso="Bold", color=S.BLANCO, interlineado=1.05)
    return s


def ficha(prs, c, plano_nuevo=None):
    """Ficha del centro: identificación + contexto satelital + planimetría."""
    s = lamina(prs, S.PAPEL)
    # cabecera
    S.rect(s, 0, 0, A, int(AL*0.195), S.ROJO)
    S.chevron(s, int(A*0.945), int(AL*0.062), int(AL*0.072), S.BLANCO, transparencia=0.72)
    S.texto(s, M, int(AL*0.045), int(A*0.55), Emu(400000), c["nombre"],
            tam=27, peso="Bold", color=S.BLANCO)
    S.texto(s, M, int(AL*0.125), int(A*0.55), Emu(300000),
            f'{c["direccion"]}, {c["comuna"]}   ›   ver en Google Maps',
            tam=11.5, peso="Light", color=S.BLANCO, enlace=maps(c))
    S.texto(s, int(A-M-Emu(2900000)), int(AL*0.068), Emu(2100000), Emu(400000),
            c["region"], tam=12, peso="SemiBold", color=S.BLANCO,
            alineado=PP_ALIGN.RIGHT, espaciado=2.0)

    # cuerpo: satelital a la izquierda, plano a la derecha
    y0, h = int(AL*0.255), int(AL*0.60)
    sats = c["assets"]["satelites"]
    if sats:
        foto_cubriendo(s, BASE/sats[0]["ruta"], M, y0, int(A*0.40), h)
        S.texto(s, M, int(y0+h+Emu(80000)), int(A*0.40), Emu(260000), "UBICACIÓN Y ENTORNO",
                tam=10, peso="SemiBold", color=S.ROJO, espaciado=2.4)
    px = int(M + A*0.40 + Emu(400000))
    plano = plano_nuevo or (BASE/c["assets"]["planos"][0]["ruta"] if c["assets"]["planos"] else None)
    if plano:
        foto_entera(s, plano, px, y0, int(A - px - M), h)
        S.texto(s, px, int(y0+h+Emu(80000)), int(A*0.4), Emu(260000), "PLANTA Y LOCALES",
                tam=10, peso="SemiBold", color=S.ROJO, espaciado=2.4)
    S.marca_pie(s)
    return s


def cierre(prs):
    s = lamina(prs)
    S.rect(s, 0, 0, A, AL, S.ROJO)
    S.campo(s, A, AL, 0.60, S.ROJO_OS, punta=0.16, transparencia=0.55)
    s.shapes.add_picture(str(BASE/"assets/logo-mascenter-blanco.png"),
                         M, int(AL*0.22), Emu(2400000))
    S.texto(s, M, int(AL*0.52), int(A*0.6), Emu(900000),
            "Conversemos sobre tu próximo local.", tam=26, peso="Bold", color=S.BLANCO)
    datos = [("Correo", "arriendos@mascenter.cl"), ("Sitio", "mascenter.cl"),
             ("Instagram", "@mascenter")]
    x = M
    for et, v in datos:
        S.texto(s, x, int(AL*0.70), Emu(2600000), Emu(250000), et.upper(),
                tam=10, peso="SemiBold", color=S.BLANCO, espaciado=2.4)
        S.texto(s, x, int(AL*0.755), Emu(2600000), Emu(300000), v,
                tam=14, peso="Light", color=S.BLANCO)
        x += Emu(2900000)
    s.shapes.add_picture(str(BASE/"assets/logo-grupoifb-blanco.png"),
                         int(A-M-Emu(1000000)), int(AL-M-Emu(620000)), Emu(1000000))
    return s


def construir(salida, muestra=True):
    cs = json.load(open(BASE/"centros.json", encoding="utf-8"))
    prs = Presentation(); prs.slide_width, prs.slide_height = A, AL
    porc = {c["nombre"]: c for c in cs}
    hero = BASE/porc["Santa María"]["assets"]["satelites"][0]["ruta"]

    portada(prs, hero)
    resumen(prs, BASE/porc["Ciudad Empresarial"]["assets"]["satelites"][0]["ruta"])
    indice(prs, [("La red Más Center", "Cobertura, escala y respaldo de Grupo IFB"),
                 ("Los centros en operación", "Ficha, ubicación y planta de cada proyecto"),
                 ("Arrienda tu local", "Cómo avanzar y con quién hablar")])
    portadilla(prs, 1, "La red\nMás Center", BASE/porc["Los Ángeles"]["assets"]["satelites"][0]["ruta"])
    ficha(prs, porc["Ciudad Empresarial"], plano_nuevo="/Users/Vale/Downloads/KV-PLANOS.png")
    ficha(prs, porc["Copiapó"])
    ficha(prs, porc["Los Ángeles"])
    cierre(prs)

    S.tema_hipervinculo(prs)          # sin esto los enlaces salen azules y subrayados
    prs.save(salida)
    return salida


if __name__ == "__main__":
    out = RAIZ/"out/mascenter-presentacion"
    out.mkdir(parents=True, exist_ok=True)
    r = construir(out/"MasCenter_Proyectos_Operativos_MUESTRA.pptx")
    print("listo:", r)
