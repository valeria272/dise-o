# -*- coding: utf-8 -*-
"""Sistema de diseño Más Center para presentaciones — medido del brochure de Algarrobal.

Todo lo de acá salió de medir el PDF del brochure (InDesign, 20 páginas), no de
inventar: la paleta se sacó contando píxeles de las páginas de color plano y el
chevron se trazó fila por fila desde el logo. Ver docs/SISTEMA-DE-MARCAS.md.
"""
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── paleta medida del brochure ───────────────────────────────────────────────
ROJO    = RGBColor(0xE4, 0x20, 0x26)   # rojo corporativo, el dominante
ROJO_OS = RGBColor(0x8F, 0x1C, 0x25)   # rojo oscuro, para profundidad
VELO    = RGBColor(0xD4, 0x39, 0x2F)   # rojo del velo sobre fotografía
GRAFITO = RGBColor(0x3C, 0x3E, 0x45)   # gris azulado de los paneles oscuros
PAPEL   = RGBColor(0xF5, 0xF3, 0xF3)   # blanco roto de los fondos
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)

FUENTE = "Poppins"                      # la del brochure, y además es Google Font

# ── lienzo 16:9 ──────────────────────────────────────────────────────────────
ANCHO, ALTO = Emu(12192000), Emu(6858000)
MARGEN = Emu(685800)                    # 0,75"

# ── el chevron: trazado del logo, fila por fila ──────────────────────────────
# Polígono de 6 vértices en coordenadas normalizadas. La punta va a media altura
# y la muesca izquierda la refleja. Proporción ancho/alto medida = 0,809.
CHEVRON = [(0.056, 0.0), (0.677, 0.0), (1.0, 0.5), (0.505, 1.0), (0.0, 1.0), (0.399, 0.5)]
CHEVRON_RATIO = 0.809


def chevron(slide, x, y, alto, color, transparencia=None, ratio=CHEVRON_RATIO):
    """Dibuja el chevron de la marca como forma nativa (editable en PowerPoint)."""
    ancho = int(alto * ratio)
    pts = [(int(x + px*ancho), int(y + py*alto)) for px, py in CHEVRON]
    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1])
    builder.add_line_segments(pts[1:], close=True)
    forma = builder.convert_to_shape()
    forma.fill.solid(); forma.fill.fore_color.rgb = color
    forma.line.fill.background()
    if transparencia is not None:
        _transparencia(forma, transparencia)
    return forma


def _transparencia(forma, pct):
    """python-pptx no expone transparencia de relleno; se escribe en el XML."""
    from pptx.oxml.ns import qn
    srgb = forma.fill.fore_color._xFill.find(qn('a:srgbClr'))
    alpha = srgb.makeelement(qn('a:alpha'), {'val': str(int((1-pct)*100000))})
    srgb.append(alpha)


def campo(slide, slide_ancho, slide_alto, frac, color, punta=0.13, transparencia=None):
    """El recurso real del brochure: el campo de color termina en punta de chevron.

    No es una flecha flotando encima de la foto — es el BORDE entre el color plano y
    la fotografía. Confundir las dos cosas fue el primer error de esta pieza.
    """
    x = int(slide_ancho*frac)
    pts = [(0,0), (x,0), (int(x+slide_ancho*punta), slide_alto//2), (x, slide_alto), (0, slide_alto)]
    b = slide.shapes.build_freeform(pts[0][0], pts[0][1])
    b.add_line_segments(pts[1:], close=True)
    f = b.convert_to_shape()
    f.fill.solid(); f.fill.fore_color.rgb = color; f.line.fill.background()
    f.shadow.inherit = False
    if transparencia is not None: _transparencia(f, transparencia)
    return f


def rect(slide, x, y, ancho, alto, color, transparencia=None):
    from pptx.enum.shapes import MSO_SHAPE
    f = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, ancho, alto)
    f.fill.solid(); f.fill.fore_color.rgb = color; f.line.fill.background()
    f.shadow.inherit = False
    if transparencia is not None: _transparencia(f, transparencia)
    return f


def pildora(slide, x, y, ancho, alto, color):
    from pptx.enum.shapes import MSO_SHAPE
    f = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, ancho, alto)
    f.fill.solid(); f.fill.fore_color.rgb = color; f.line.fill.background()
    f.shadow.inherit = False
    f.adjustments[0] = 0.5
    return f


def texto(slide, x, y, ancho, alto, contenido, *, tam=18, peso="Regular", color=GRAFITO,
          alineado=PP_ALIGN.LEFT, interlineado=1.15, espaciado=0, anclaje=MSO_ANCHOR.TOP,
          enlace=None, mayusculas=False):
    """Caja de texto NATIVA — editable en PowerPoint y en Google Slides."""
    caja = slide.shapes.add_textbox(x, y, ancho, alto)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anclaje
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lineas = contenido.split("\n") if isinstance(contenido, str) else contenido
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alineado
        p.line_spacing = interlineado
        r = p.add_run()
        r.text = linea.upper() if mayusculas else linea
        fnt = r.font
        # Google Slides sólo reconoce la FAMILIA: "Poppins SemiBold" no existe para
        # él y cae a Arial. El grosor se pide con bold, no con el nombre.
        fnt.name = f"{FUENTE} Light" if peso == "Light" else FUENTE
        fnt.size = Pt(tam); fnt.color.rgb = color
        fnt.bold = peso in ("Bold", "SemiBold")
        if espaciado:
            from pptx.oxml.ns import qn
            r._r.get_or_add_rPr().set('spc', str(int(espaciado*100)))
        if enlace:
            r.hyperlink.address = enlace
            r._r.get_or_add_rPr().set('u', 'none')   # el color lo fija el tema:
                                                     # ver tema_hipervinculo()
    return caja


def marca_pie(slide, oscuro=False):
    """El chevron chico abajo a la derecha — va en todas las páginas del brochure."""
    alto = Emu(200000)
    chevron(slide, int(ANCHO - MARGEN - alto*CHEVRON_RATIO), int(ALTO - MARGEN*0.55 - alto),
            alto, ROJO if not oscuro else BLANCO)


def tema_hipervinculo(prs, color="FFFFFF"):
    """Pinta los hipervínculos del color pedido cambiando el TEMA.

    PowerPoint ignora el relleno del run para los enlaces y usa el color `hlink`
    del esquema del tema; por eso el azul subrayado sobrevivía a todo intento de
    pintarlo run por run. Se cambia una vez y vale para toda la presentación.
    """
    from pptx.oxml.ns import qn
    from lxml import etree
    for parte in prs.part.package.iter_parts():
        if "theme" not in str(parte.partname): continue
        arbol = etree.fromstring(parte.blob)
        esquema = arbol.find(qn('a:themeElements') + '/' + qn('a:clrScheme'))
        if esquema is None: continue
        for etiqueta in ('a:hlink', 'a:folHlink'):
            nodo = esquema.find(qn(etiqueta))
            if nodo is None: continue
            for hijo in list(nodo): nodo.remove(hijo)
            nodo.append(nodo.makeelement(qn('a:srgbClr'), {'val': color}))
        parte._blob = etree.tostring(arbol, xml_declaration=True,
                                     encoding='UTF-8', standalone=True)
        return True
    return False
